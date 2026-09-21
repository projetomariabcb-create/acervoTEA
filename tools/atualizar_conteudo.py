from pathlib import Path
from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from html import escape
import re, json, shutil, unicodedata

ROOT=Path(__file__).resolve().parents[1]
SOURCE=next((ROOT/'fonte').glob('*.docx'), None)
if SOURCE is None: raise SystemExit('Nenhum .docx encontrado em fonte/.')

def iter_blocks(parent):
    elm=parent.element.body if isinstance(parent,_Document) else parent._tc
    for child in elm.iterchildren():
        if isinstance(child,CT_P): yield Paragraph(child,parent)
        elif isinstance(child,CT_Tbl): yield Table(child,parent)
url_re=re.compile(r'(https?://[^\s<>()]+)'); ref_re=re.compile(r'\[(R\d{2,3})\]')
def linkify(s):
    e=escape(s)
    e=url_re.sub(lambda m:f'<a href="{m.group(1)}" target="_blank" rel="noopener noreferrer">{m.group(1)}</a>',e)
    e=ref_re.sub(lambda m:f'<a class="ref-link" href="#ref-{m.group(1).lower()}" title="Ir para {m.group(1)}">[{m.group(1)}]</a>',e)
    return e
def runs(p):
    if not p.runs:return linkify(p.text)
    o=[]
    for r in p.runs:
        t=linkify(r.text)
        if r.bold:t=f'<strong>{t}</strong>'
        if r.italic:t=f'<em>{t}</em>'
        if r.underline:t=f'<u>{t}</u>'
        if r.font.superscript:t=f'<sup>{t}</sup>'
        if r.font.subscript:t=f'<sub>{t}</sub>'
        o.append(t)
    return ''.join(o)
def slug(s):
    s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',s).strip('-')[:90] or 'secao'
seen={}
def uid(s):
    b=slug(s); seen[b]=seen.get(b,0)+1; return b if seen[b]==1 else f'{b}-{seen[b]}'
def table_html(t):
    heads=[c.text.strip() for c in t.rows[0].cells] if t.rows else []; refs=heads[:2]==['ID','Referência/fonte']; a=['<div class="table-wrap"><table>']
    for ri,row in enumerate(t.rows):
        tag='th' if ri==0 else 'td'; a.append('<tr>')
        for ci,c in enumerate(row.cells):
            attr=''
            if refs and ri and ci==0:attr=f' id="ref-{escape(c.text.strip().lower())}"'
            a.append(f'<{tag}{attr}>'+ '<br>'.join(linkify(x) for x in c.text.split('\n'))+f'</{tag}>')
        a.append('</tr>')
    a.append('</table></div>'); return ''.join(a)

doc=Document(SOURCE); html=[]; toc=[]; current_part=None; list_open=False

def close_list():
    global list_open
    if list_open: html.append('</ul>'); list_open=False
for b in iter_blocks(doc):
    if isinstance(b,Paragraph):
        txt=b.text.strip(); st=b.style.name if b.style else ''
        if not txt: close_list(); continue
        if st=='Heading 1':
            close_list(); i=uid(txt); current_part=i; toc.append({'level':1,'id':i,'title':txt}); html.extend([f'<section class="part" id="{i}" data-search-title="{escape(txt)}">',f'<div class="part-kicker">Capítulo</div><h1>{linkify(txt)}</h1>','</section>'])
        elif st=='Heading 2':
            close_list(); i=uid(txt); toc.append({'level':2,'id':i,'title':txt,'part':current_part}); html.extend([f'<section class="topic" id="{i}" data-search-title="{escape(txt)}"><h2>{linkify(txt)}</h2>','</section>'])
        elif st=='List Bullet':
            if not list_open:html.append('<ul class="content-list">');list_open=True
            html.append(f'<li>{runs(b)}</li>')
        elif st=='Callout':close_list();html.append(f'<aside class="callout"><span class="callout-icon" aria-hidden="true">◆</span><div>{runs(b)}</div></aside>')
        elif st=='Lead':close_list();html.append(f'<p class="lead">{runs(b)}</p>')
        elif st=='Small':close_list();html.append(f'<p class="small-note">{runs(b)}</p>')
        else:close_list();html.append(f'<p>{runs(b)}</p>')
    else:close_list();html.append(table_html(b))
close_list()
search=[]; it=iter(toc); current=next(it,None); cur_id='topo'; cur_title='Introdução'; cur_part=''; used=0
# second pass pairing headings sequentially
headings=[x for x in toc]
hi=0
for b in iter_blocks(doc):
    if isinstance(b,Paragraph):
        txt=b.text.strip(); st=b.style.name if b.style else ''
        if st in ('Heading 1','Heading 2') and txt:
            item=headings[hi];hi+=1;cur_id=item['id'];cur_title=txt
            if st=='Heading 1':cur_part=txt
            search.append({'id':cur_id,'title':cur_title,'part':cur_part,'text':''})
        elif txt:
            if not search:search.append({'id':'topo','title':'Introdução','part':'','text':''})
            search[-1]['text']+=' '+txt
    elif isinstance(b,Table) and search:search[-1]['text']+=' '+' '.join(c.text for r in b.rows for c in r.cells)
for s in search:s['text']=re.sub(r'\s+',' ',s['text']).strip()
updated='21 de setembro de 2026'
for p in doc.paragraphs[:20]:
    m=re.search(r'Atualizado até\s+(.+?)(?:\.|$)',p.text,re.I)
    if m: updated=m.group(1).strip()
refs=0
for t in doc.tables:
    if t.rows and [c.text.strip() for c in t.rows[0].cells][:2]==['ID','Referência/fonte']: refs=max(0,len(t.rows)-1)
meta={'title':'Dossiê Científico do Autismo','subtitle':'Estado da Arte, História, Evidências, Intervenções, Direitos e Perspectivas','period':'1911–2026','updated':updated,'description':'Acervo informativo baseado em evidências sobre Transtorno do Espectro Autista (TEA), para famílias, profissionais, educadores e pesquisadores.','parts':sum(1 for x in toc if x['level']==1 and x['title'].startswith('PARTE')),'topics':sum(1 for x in toc if x['level']==2),'references':refs,'sourcePages':52}
out='window.DOSSIE = '+json.dumps({'meta':meta,'toc':toc,'search':search,'html':'\n'.join(html)},ensure_ascii=False)+';\n'
(ROOT/'content/conteudo.js').write_text(out,encoding='utf-8')

# Regenera também o glossário estruturado de siglas/abreviações/símbolos para tooltips do site.
acro={}
for t in doc.tables:
    if t.rows and [c.text.strip() for c in t.rows[0].cells][:4]==['Categoria','Sigla / símbolo','Nome por extenso','Contexto / observação']:
        for row in t.rows[1:]:
            cat,sig,full,obs=[c.text.strip() for c in row.cells]
            if sig:
                acro[sig]={'full':full,'context':obs,'category':cat}
        break
(ROOT/'content'/'siglas.js').write_text('window.SIGLAS = '+json.dumps(acro,ensure_ascii=False)+';\n',encoding='utf-8')

shutil.copy2(SOURCE,ROOT/'downloads'/SOURCE.name)
print(f'Conteúdo atualizado: {len(toc)} entradas de navegação, {refs} referências.')
print('Observação: o PDF em downloads/ não é regenerado por este script; exporte um novo PDF do Word se necessário.')
